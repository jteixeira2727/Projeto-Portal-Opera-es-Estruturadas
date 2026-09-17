#!/usr/bin/env python3
"""
Extrai texto dos documentos de cada operação (pasta Operações/<Op>/Documentos/
staged em /mnt/user-data/uploads/...) e grava documents_data.json:
  { "<OperationName>": [ {"filename": ..., "text": ...}, ... ], ... }

Roda uma vez (ou sempre que a pasta de documentos mudar) — separado do
build_data.py, que roda a cada rodada de ajuste.
"""
import os, sys, json, subprocess, re

UPLOADS_ROOT = "/mnt/user-data/uploads/Projeto Portal Operações Estruturadas/Operações"
OUT_PATH = "/tmp/work/documents_data.json"

# OCR fallback (achado ao investigar a reclamação do usuário "pergunto e ele não me
# traz resposta"): em pelo menos 3 das 18 atas da Alianza (AGT 04.08.2022, 07.11.2023,
# 14.03.2022), o PDF é uma assinatura via DocuSign/certificado onde a página inteira
# foi "achatada" numa imagem (confirmado via PyMuPDF: text_len=0, 1+ imagem por página)
# -- nem pdftotext nem pdfplumber extraem nada disso, e o doc_chunk final fica
# permanentemente vazio pra essas atas, mesmo com o arquivo presente na pasta. Fix:
# processa por PÁGINA (não por arquivo inteiro) via PyMuPDF; quando uma página
# individual não tem texto nativo, renderiza em imagem (200dpi) e roda OCR (tesseract,
# idioma português) só nela -- preserva as páginas que já tinham texto nativo (mais
# rápidas e mais precisas) e só paga o custo de OCR onde é realmente necessário.
def extract_pdf_text(path):
    try:
        import fitz  # PyMuPDF
    except Exception as e:
        print(f"  PyMuPDF indisponível ({e}), caindo pro pdftotext simples...", file=sys.stderr)
        try:
            out = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, timeout=120)
            return out.stdout.decode("utf-8", errors="replace")
        except Exception as e2:
            print(f"  pdftotext também falhou ({e2})", file=sys.stderr)
            return ""

    # Limiar de 300 chars (2ª correção, achado ao investigar BR Properties): algumas
    # páginas de escritura escaneada (ex.: "AF Galpão G100-F.pdf", uma matrícula de
    # imóvel de 2 páginas) têm uma FINA camada de texto nativo sobreposta -- só o link
    # de validação de assinatura digital ("Valide este documento clicando no link a
    # seguir...", ~138 chars) -- por cima de uma imagem com o conteúdo real da página
    # (a matrícula em si). Confirmado (276 chars extraídos no total pras 2 páginas,
    # contra >2000 chars/página que o OCR recupera) que essas páginas TÊM `t.strip()`
    # não-vazio, então o check anterior (`if t.strip()`) as tratava como "já tem texto"
    # e nunca tentava OCR -- perdendo o conteúdo real (nesse caso, dados de registro de
    # garantia). Fix: só aceita o texto nativo direto quando ele for "denso" o
    # suficiente (>=300 chars) OU quando a página não tem nenhuma imagem embutida (texto
    # nativo puro, sem risco de estar sobre um scan) -- caso contrário, tenta OCR e,
    # se vier algo, usa o OCR (mais completo) em vez do texto nativo raso.
    MIN_NATIVE_CHARS = 300
    parts = []
    ocr_pages = 0
    try:
        doc = fitz.open(path)
        for page in doc:
            t = page.get_text()
            has_images = bool(page.get_images())
            if len(t.strip()) >= MIN_NATIVE_CHARS or (t.strip() and not has_images):
                parts.append(t)
                continue
            # texto nativo ausente ou raso demais numa página com imagem -- tenta OCR
            try:
                import pytesseract
                from PIL import Image
                import io
                pix = page.get_pixmap(dpi=200)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                ocr_t = pytesseract.image_to_string(img, lang='por')
                if ocr_t.strip():
                    parts.append(ocr_t)
                    ocr_pages += 1
                elif t.strip():
                    parts.append(t)  # OCR não achou nada -- melhor o pouco nativo que nada
            except Exception as e:
                print(f"  OCR falhou numa página ({e})", file=sys.stderr)
                if t.strip():
                    parts.append(t)
        doc.close()
    except Exception as e:
        print(f"  PyMuPDF falhou ({e}), tentando pdftotext/pdfplumber...", file=sys.stderr)
        try:
            out = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, timeout=120)
            text = out.stdout.decode("utf-8", errors="replace")
            if text.strip():
                return text
        except Exception:
            pass
        try:
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                return "\n".join(p.extract_text() or "" for p in pdf.pages)
        except Exception as e2:
            print(f"  pdfplumber também falhou ({e2})", file=sys.stderr)
            return ""

    if ocr_pages:
        print(f"  ({ocr_pages} página(s) via OCR)")
    return "\n".join(parts)


def extract_docx_text(path):
    import docx
    d = docx.Document(path)
    parts = []
    for p in d.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def clean_text(t):
    # normaliza espaços em excesso mas preserva quebras de parágrafo
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def main():
    if not os.path.isdir(UPLOADS_ROOT):
        print(f"Pasta não encontrada: {UPLOADS_ROOT}", file=sys.stderr)
        sys.exit(1)

    result = {}
    for op_name in sorted(os.listdir(UPLOADS_ROOT)):
        doc_dir = os.path.join(UPLOADS_ROOT, op_name, "Documentos")
        if not os.path.isdir(doc_dir):
            continue
        docs = []
        for fname in sorted(os.listdir(doc_dir)):
            fpath = os.path.join(doc_dir, fname)
            if not os.path.isfile(fpath):
                continue
            ext = os.path.splitext(fname)[1].lower()
            print(f"[{op_name}] extraindo: {fname} ({os.path.getsize(fpath)} bytes)")
            if ext == ".pdf":
                text = extract_pdf_text(fpath)
            elif ext == ".docx":
                text = extract_docx_text(fpath)
            else:
                print(f"  tipo não suportado ({ext}), pulando")
                continue
            text = clean_text(text)
            print(f"  -> {len(text)} caracteres extraídos")
            docs.append({"filename": fname, "text": text})
        if docs:
            result[op_name] = docs

    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False)
    total_chars = sum(len(d["text"]) for docs in result.values() for d in docs)
    print(f"\nEscrito {OUT_PATH}: {len(result)} operação(ões), {sum(len(v) for v in result.values())} documento(s), {total_chars} caracteres totais")


if __name__ == "__main__":
    main()
